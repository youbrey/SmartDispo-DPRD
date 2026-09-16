import shutil
import subprocess
from datetime import date, datetime
from pathlib import Path
from uuid import UUID

from docx import Document as WordDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models.entities import (
    Disposition,
    DispositionSheet,
    DispositionTarget,
    Document,
    DocumentTemplateVersion,
    DocumentVersion,
    OrganizationalUnit,
    Role,
    RoleAssignment,
    User,
)
from app.services.templates import resolve_template_path

TEMPLATES = {
    "MEETING_REQUEST": "meeting-request-v1.docx",
    "TRAVEL_REQUEST": "travel-request-v1.docx",
    "INCOMING_CHAIRMAN": "disposition-dprd-v1.docx",
    "INCOMING_SECRETARY": "disposition-setwan-v1.docx",
}


def _date(value: str | date | datetime | None) -> str:
    if not value:
        return ""
    parsed = value if isinstance(value, (date, datetime)) else date.fromisoformat(str(value)[:10])
    months = [
        "Januari",
        "Februari",
        "Maret",
        "April",
        "Mei",
        "Juni",
        "Juli",
        "Agustus",
        "September",
        "Oktober",
        "November",
        "Desember",
    ]
    return f"{parsed.day} {months[parsed.month - 1]} {parsed.year}"


def _set_cell(table, row: int, column: int, value: str) -> None:
    cell = table.rows[row].cells[column]
    paragraph = cell.paragraphs[0]
    _set_paragraph(paragraph, value)
    for extra in cell.paragraphs[1:]:
        _set_paragraph(extra, "")


def _set_paragraph(paragraph, value: str, size: float | None = None) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    if size is not None:
        for run in paragraph.runs:
            run.font.size = Pt(size)


def _insert_paragraph_after(paragraph: Paragraph, value: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    created = Paragraph(element, paragraph._parent)
    _set_paragraph(created, value)
    return created


def _number_word(value: int) -> str:
    units = [
        "nol",
        "satu",
        "dua",
        "tiga",
        "empat",
        "lima",
        "enam",
        "tujuh",
        "delapan",
        "sembilan",
        "sepuluh",
        "sebelas",
    ]
    if value < len(units):
        return units[value]
    if value < 20:
        return f"{units[value - 10]} belas"
    if value < 100:
        tens, remainder = divmod(value, 10)
        return f"{units[tens]} puluh" + (f" {units[remainder]}" if remainder else "")
    return str(value)


def _display_person(person: dict | None) -> str:
    if not person:
        return "................................................"
    metadata = person.get("metadata") or {}
    name = person.get("full_name") or ""
    nip = metadata.get("nip")
    rank = metadata.get("rank") or metadata.get("pangkat_golongan")
    details = [name.upper()]
    if rank:
        details.append(str(rank).upper())
    if nip:
        details.append(f"NIP. {nip}")
    return "\n".join(details)


def _format_table(table, size: float = 9.5) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(size)


def _rebuild_setwan_sheet(
    word: WordDocument,
    content: dict,
    directives: set[str],
    notes: list[str],
    officeholders: dict,
    target_labels: list[str],
) -> None:
    old_table = word.tables[0]
    old_table._element.getparent().remove(old_table._element)
    table = word.add_table(rows=9, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(3.35)
    table.columns[1].width = Inches(3.35)

    title = table.cell(0, 0).merge(table.cell(0, 1))
    _set_paragraph(title.paragraphs[0], "LEMBAR DISPOSISI", 14)
    title.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraphs[0].runs[0].bold = True
    _set_cell(table, 1, 0, f"Surat Dari : {content['sender']}")
    _set_cell(table, 1, 1, f"Diterima Tgl. : {_date(content['received_date'])}")
    _set_cell(table, 2, 0, f"No. Surat : {content['letter_number']}")
    _set_cell(table, 2, 1, f"No. Agenda : {content['agenda_number']}")

    priority = content.get("priority", "BIASA")
    priority_line = "    ".join(
        f"{'☒' if priority == code else '☐'} {label}"
        for code, label in (
            ("BIASA", "Biasa"),
            ("PENTING", "Penting"),
            ("SEGERA", "Segera"),
            ("RAHASIA", "Rahasia"),
        )
    )
    _set_cell(table, 3, 0, f"Tanggal : {_date(content['letter_date'])}\nSifat : {priority_line}")
    _set_cell(table, 3, 1, "")
    subject = table.cell(4, 0).merge(table.cell(4, 1))
    _set_paragraph(subject.paragraphs[0], f"Perihal : {content['subject']}")
    coordination = table.cell(5, 0).merge(table.cell(5, 1))
    _set_paragraph(coordination.paragraphs[0], "Paraf Koordinasi : Kabag Umum & Keuangan\n\n")

    targets = "Diteruskan Kepada :\n" + (
        "\n".join(f"☒ {label}" for label in target_labels)
        if target_labels
        else "☐ Belum ada tujuan disposisi"
    )
    setwan_labels = {
        "FURTHER_PROCESS": "Proses Lebih Lanjut",
        "CREATE_REVIEW_ADVICE": "Buat Telaahan dan Saran",
        "COORDINATE": "Koordinasikan",
        "STUDY_REPORT": "Pelajari dan Laporkan",
        "MONITOR_INPUT": "Monitor Untuk Masukan",
        "CONSIDER": "Pertimbangkan",
        "GUIDANCE": "Untuk Dipedomani",
        "PREPARE_MATERIAL": "Buat Materi/Siapkan Bahan",
        "ATTENTION": "Untuk Minta Perhatian",
        "ACKNOWLEDGE": "Untuk Diketahui",
        "CREATE_SPT": "Buatkan SPT",
        "CREATE_SPD": "Buatkan SPD",
        "FILE": "File",
    }
    directive_text = "Mengharapkan :\n" + "\n".join(
        f"{'☒' if code in directives else '☐'} {label}" for code, label in setwan_labels.items()
    )
    _set_cell(table, 6, 0, targets)
    _set_cell(table, 6, 1, directive_text)
    _set_cell(table, 7, 0, f"CATATAN SEKWAN :\n{notes[-1] if notes else ''}\n\n")
    _set_cell(table, 7, 1, "Catatan Kabag :\n\n")
    signature = table.cell(8, 0).merge(table.cell(8, 1))
    _set_paragraph(
        signature.paragraphs[0],
        "SEKRETARIS DPRD KOTA BITUNG,\n\n" + _display_person(officeholders.get("SEKWAN")),
    )
    signature.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_table(table)


def _fill_meeting(word: WordDocument, content: dict) -> None:
    purpose_lines = [p for p in word.paragraphs if p.text.count(".") > 80]
    _set_cell(word.tables[0], 0, 2, content["sender_name"])
    _set_cell(word.tables[0], 1, 2, content["sender_position"])
    scheduled = datetime.fromisoformat(content["scheduled_at"])
    _set_cell(word.tables[1], 0, 3, _date(scheduled))
    _set_cell(word.tables[1], 1, 3, scheduled.strftime("%H.%M WITA"))
    _set_cell(word.tables[1], 2, 3, content["place"])
    _set_cell(word.tables[1], 3, 3, content.get("attire") or "-")
    for paragraph in word.paragraphs:
        if "Akan mengadakan @Jenis_rapat" in paragraph.text:
            _set_paragraph(
                paragraph,
                (
                    "Bersama ini disampaikan kepada Pimpinan DPRD Kota Bitung, bahwa Pimpinan dan Anggota "
                    f"DPRD Kota Bitung akan mengadakan {content['meeting_type_name']}, dalam rangka :"
                ),
            )
        elif paragraph.text.strip().startswith("Catatan"):
            _set_paragraph(paragraph, f"Catatan : {content.get('notes') or '-'}")
    purpose_slot = purpose_lines[0] if purpose_lines else None
    if purpose_slot:
        _set_paragraph(purpose_slot, content["purpose"])
    for paragraph in purpose_lines[1:]:
        _set_paragraph(paragraph, "")

    invitees = [
        f"{index}. {item['name']}" + (f" - {item['institution']}" if item.get("institution") else "")
        for index, item in enumerate(content["invitees"], 1)
    ]
    invitation_table = word.tables[2].rows[0].cells[0].tables[0]
    for row in range(10):
        left = invitees[row] if row < len(invitees) else f"{row + 1}."
        right_index = row + 10
        right = invitees[right_index] if right_index < len(invitees) else f"{right_index + 1}."
        _set_cell(invitation_table, row, 0, left)
        _set_cell(invitation_table, row, 1, right)
    _format_table(invitation_table, 8.5)


def _fill_travel(word: WordDocument, content: dict) -> None:
    _set_cell(word.tables[0], 0, 2, content["sender_name"])
    _set_cell(word.tables[5], 0, 2, content["sender_name"])
    _set_cell(word.tables[5], 1, 2, content["sender_position"])
    _set_cell(word.tables[6], 0, 2, f"{_date(content['start_date'])} s.d. {_date(content['end_date'])}")
    _set_cell(word.tables[6], 1, 2, content.get("activity_time") or "-")
    _set_cell(word.tables[6], 2, 2, content["place"])
    executors = [member for member in content["members"] if member["member_group"] == "EXECUTOR"]
    companions = [member for member in content["members"] if member["member_group"] == "ACCOMPANYING"]
    for index in range(12):
        row = index % 6 + 1
        column = 0 if index < 6 else 1
        name = executors[index]["name"] if index < len(executors) else ""
        _set_cell(word.tables[1], row, column, f"{index + 1}. {name}" if name else f"{index + 1}.")
    for row in range(1, 14):
        _set_cell(word.tables[7], row, 0, f"{row}. {executors[row - 1]['name']}" if row <= len(executors) else "")
        _set_cell(word.tables[7], row, 1, f"{row}. {companions[row - 1]['name']}" if row <= len(companions) else "")
    follow_up_labels = {
        "WORK_MEETING": "Rapat kerja",
        "HEARING": "Rapat dengar pendapat",
        "COORDINATE": "Koordinasikan",
        "CONSULT": "Konsultasikan",
        "RECOMMEND": "Rekomendasikan",
        "ARCHIVE": "Arsip",
    }
    selected_follow_ups = set(content.get("follow_up_directives") or [])
    for row, (code, label) in enumerate(follow_up_labels.items()):
        target_row = 8 if code == "ARCHIVE" else row
        _set_cell(word.tables[4], target_row, 0, "☒" if code in selected_follow_ups else "☐")
        _set_cell(word.tables[4], target_row, 1, label)
    destinations = list(content["destinations"])
    start = date.fromisoformat(str(content["start_date"])[:10])
    end = date.fromisoformat(str(content["end_date"])[:10])
    duration = (end - start).days + 1
    dotted_slots: dict[str, list] = {
        "destination": [],
        "material": [],
        "problem": [],
        "condition": [],
        "efforts": [],
        "purpose": [],
    }
    section = ""
    for paragraph in word.paragraphs:
        text = paragraph.text.strip()
        if text == "KE :":
            section = "destination"
        elif text.startswith("MATERI KONSULTASI"):
            section = "material"
        elif text.startswith("1.   Permasalahan umum"):
            section = "problem"
        elif text.startswith("2.   Kondisi saat ini"):
            section = "condition"
        elif text.startswith("3.   Upaya yang telah"):
            section = "efforts"
        elif text.startswith("Waktu pelaksanaan"):
            _set_paragraph(paragraph, f"Waktu pelaksanaan\t:\t{duration} ({_number_word(duration)}) Hari")
            section = ""
        elif text.startswith("Terhitung"):
            _set_paragraph(paragraph, f"Terhitung\t:\t{_date(start)} s/d {_date(end)}")
            section = ""
        elif text.startswith("Bersama ini disampaikan"):
            activity = "Konsultasi" if content["activity_type"] == "CONSULTATION" else "Kunjungan Kerja"
            _set_paragraph(
                paragraph,
                f"Bersama ini disampaikan kepada Pimpinan DPRD Kota Bitung, bahwa {content['organizational_unit']} "
                f"akan mengadakan {activity}, dalam rangka / tentang:",
            )
            section = "purpose"
        elif text.startswith("yang akan dilaksanakan"):
            section = ""
        elif text.count(".") > 80 and section:
            dotted_slots[section].append(paragraph)

    values = {
        "destination": destinations,
        "material": [content["material"]],
        "problem": [content["general_problem"]],
        "condition": [content["current_condition"]],
        "efforts": [content["efforts"]],
        "purpose": [content["purpose"]],
    }
    for key, slots in dotted_slots.items():
        for index, paragraph in enumerate(slots):
            value = values[key][index] if index < len(values[key]) else ""
            prefix = f"{index + 1}. " if key == "destination" else ""
            _set_paragraph(paragraph, f"{prefix}{value}" if value else "")
    for paragraph in word.paragraphs[-3:]:
        if not paragraph.text.strip():
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            if paragraph.runs:
                paragraph.runs[0].font.size = Pt(1)


def _fill_disposition(word: WordDocument, content: dict, kind: str) -> None:
    directives = {
        directive for entry in (content.get("dispositions") or {}).values() for directive in entry.get("directives", [])
    }
    notes = [entry.get("note") for entry in (content.get("dispositions") or {}).values() if entry.get("note")]
    officeholders = content.get("officeholders") or {}
    if kind == "INCOMING_SECRETARY":
        _rebuild_setwan_sheet(
            word,
            content,
            directives,
            notes,
            officeholders,
            (content.get("disposition_target_labels_by_actor") or {}).get("SEKWAN", []),
        )
    else:
        table = word.tables[0]
        table.rows[0].cells[0].text = f"SURAT DARI : {content['sender']}"
        table.rows[1].cells[
            0
        ].text = f"NOMOR SURAT : {content['letter_number']}    TERIMA TANGGAL : {_date(content['received_date'])}"
        table.rows[2].cells[
            0
        ].text = f"TANGGAL SURAT : {_date(content['letter_date'])}    NOMOR AGENDA : {content['agenda_number']}"
        table.rows[3].cells[0].text = f"PERIHAL : {content['subject']}"
        _set_cell(table, 9, 0, _display_person(officeholders.get("VICE_CHAIR_1")))
        _set_cell(table, 10, 0, _display_person(officeholders.get("VICE_CHAIR_2")))
        _set_cell(table, 11, 0, _display_person(officeholders.get("CHAIRMAN")))
        if notes:
            _set_cell(table, 8, 1, f"Catatan :\n{notes[-1]}")
        dprd_map = [
            [
                ("FORWARD_COMMISSION_I", "TERUSKAN KE KOMISI I"),
                ("FORWARD_COMMISSION_II", "TERUSKAN KE KOMISI II"),
                ("FORWARD_COMMISSION_III", "TERUSKAN KE KOMISI III"),
            ],
            [("FORWARD_BAPEMPERDA", "TERUSKAN KE BAPEMPERDA")],
            [("FORWARD_BANGGAR", "TERUSKAN KE BADAN ANGGARAN")],
            [("FORWARD_PANSUS", "TERUSKAN KE PANSUS")],
            [("FOLLOW_UP", "UNTUK DITINDAKLANJUTI")],
            [("ACKNOWLEDGE", "UNTUK DIKETAHUI")],
            [("REMIND", "DIINGATKAN")],
            [("POSTPONE_CANCEL", "DITUNDA/DIBATALKAN")],
        ]
        other_columns = [
            [
                "ARCHIVE",
                "CREATE_SPT",
                "CREATE_RECOMMENDATION",
                "CREATE_APPROVAL",
                "CREATE_SPEECH",
                "STUDY_RESEARCH",
                "APPROVED",
                "SCHEDULE",
            ],
            [
                "REPRESENTED_BY",
                "PROCESS_BY_MECHANISM",
                "ADJUST_BUDGET",
                "COORDINATE_CONFIRM",
                "COPY_MULTIPLY",
                "CREATE_INVITATION",
                "CREATE_DESTINATION_NOTICE",
                None,
            ],
        ]
        labels = {
            "ARCHIVE": "DIARSIPKAN",
            "CREATE_SPT": "BUATKAN SPT / SPD",
            "CREATE_RECOMMENDATION": "BUATKAN REKOMENDASI",
            "CREATE_APPROVAL": "BUATKAN PERSETUJUAN",
            "CREATE_SPEECH": "BUATKAN SAMBUTAN",
            "STUDY_RESEARCH": "DIPELAJARI/DITELITI",
            "APPROVED": "DISETUJUI",
            "SCHEDULE": "DIJADWALKAN",
            "REPRESENTED_BY": "DIWAKILI OLEH",
            "PROCESS_BY_MECHANISM": "PROSES SESUAI MEKANISME",
            "ADJUST_BUDGET": "SESUAIKAN DENGAN ANGGARAN",
            "COORDINATE_CONFIRM": "KOORDINASI/KONFIRMASI",
            "COPY_MULTIPLY": "FOTOCOPY PERBANYAK",
            "CREATE_INVITATION": "BUATKAN UNDANGAN",
            "CREATE_DESTINATION_NOTICE": "BUATKAN SURAT PEMBERITAHUAN KE DAERAH TUJUAN",
        }
        choices = word.tables[1]
        for row, options in enumerate(dprd_map):
            selected = [label for code, label in options if code in directives]
            label = selected[0] if selected else " / ".join(label for _, label in options)
            _set_cell(choices, row, 0, f"{'☒' if selected else '☐'} {label}")
        for column, codes in enumerate(other_columns, start=1):
            for row, code in enumerate(codes):
                if code:
                    selected = code in directives or (code == "CREATE_SPT" and "CREATE_SPD" in directives)
                    _set_cell(choices, row, column, f"{'☒' if selected else '☐'} {labels[code]}")
        sekwan_targets = (content.get("disposition_target_labels_by_actor") or {}).get("SEKWAN", [])
        target_paragraphs = [
            paragraph
            for paragraph in word.paragraphs
            if paragraph.text.strip().upper().startswith("KABAG ")
        ]
        for index, paragraph in enumerate(target_paragraphs):
            label = sekwan_targets[index] if index < len(sekwan_targets) else ""
            _set_paragraph(paragraph, f"☒ {label}" if label else "")
        anchor = target_paragraphs[-1] if target_paragraphs else None
        for label in sekwan_targets[len(target_paragraphs) :]:
            if anchor:
                anchor = _insert_paragraph_after(anchor, f"☒ {label}")


def render_docx(
    document: Document,
    content: dict,
    template_version: DocumentTemplateVersion | None = None,
) -> Path:
    settings = get_settings()
    template = resolve_template_path(template_version, document.document_type)
    if not template.exists():
        raise HTTPException(status_code=500, detail="Template dokumen tidak ditemukan")
    output_dir = Path(settings.generated_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"{document.id}-v{document.current_version}.docx"
    word = WordDocument(template)
    if document.document_type.value == "MEETING_REQUEST":
        _fill_meeting(word, content)
    elif document.document_type.value == "TRAVEL_REQUEST":
        _fill_travel(word, content)
    else:
        _fill_disposition(word, content, document.document_type.value)
    word.save(output)
    return output


def convert_pdf(docx_path: Path) -> Path:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise HTTPException(status_code=503, detail="Konverter PDF belum tersedia")
    subprocess.run(
        [executable, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        check=True,
        timeout=60,
        capture_output=True,
    )
    pdf = docx_path.with_suffix(".pdf")
    if not pdf.exists():
        raise HTTPException(status_code=500, detail="PDF gagal dibuat")
    return pdf


async def generate_document(session: AsyncSession, document_id: UUID, pdf: bool = False) -> Path:
    document = await session.get(Document, document_id)
    if not document:
        raise HTTPException(status_code=404, detail="Dokumen tidak ditemukan")
    version = (
        await session.execute(
            select(DocumentVersion).where(
                DocumentVersion.document_id == document.id,
                DocumentVersion.version_number == document.current_version,
            )
        )
    ).scalar_one()
    content = dict(version.content)
    template_version = (
        await session.get(DocumentTemplateVersion, version.template_version_id) if version.template_version_id else None
    )
    if document.document_type.value in {"INCOMING_CHAIRMAN", "INCOMING_SECRETARY"}:
        sheet = (
            await session.execute(select(DispositionSheet).where(DispositionSheet.document_id == document.id))
        ).scalar_one_or_none()
        content["dispositions"] = sheet.structured_data if sheet else {}
        today = date.today()
        officeholder_rows = (
            await session.execute(
                select(Role.code, User.full_name, RoleAssignment.metadata_)
                .join(RoleAssignment, RoleAssignment.role_id == Role.id)
                .join(User, User.id == RoleAssignment.user_id)
                .where(
                    Role.code.in_(["CHAIRMAN", "VICE_CHAIR_1", "VICE_CHAIR_2", "SEKWAN"]),
                    RoleAssignment.active.is_(True),
                    RoleAssignment.valid_from <= today,
                    (RoleAssignment.valid_until.is_(None) | (RoleAssignment.valid_until >= today)),
                    User.active.is_(True),
                )
            )
        ).all()
        content["officeholders"] = {
            code: {"full_name": full_name, "metadata": metadata or {}}
            for code, full_name, metadata in officeholder_rows
        }
        target_rows = (
            await session.execute(
                select(
                    Disposition.actor_role,
                    OrganizationalUnit.name,
                    Role.name,
                    User.full_name,
                )
                .select_from(DispositionTarget)
                .join(Disposition, Disposition.id == DispositionTarget.disposition_id)
                .join(DispositionSheet, DispositionSheet.id == Disposition.sheet_id)
                .outerjoin(OrganizationalUnit, OrganizationalUnit.id == DispositionTarget.unit_id)
                .outerjoin(Role, Role.id == DispositionTarget.role_id)
                .outerjoin(User, User.id == DispositionTarget.user_id)
                .where(DispositionSheet.document_id == document.id)
                .order_by(Disposition.created_at, DispositionTarget.id)
            )
        ).all()
        labels_by_actor: dict[str, list[str]] = {}
        for actor_role, unit_name, role_name, user_name in target_rows:
            label = unit_name or role_name or user_name
            if label and label not in labels_by_actor.setdefault(actor_role, []):
                labels_by_actor[actor_role].append(label)
        content["disposition_target_labels_by_actor"] = labels_by_actor
    docx = render_docx(document, content, template_version)
    return convert_pdf(docx) if pdf else docx

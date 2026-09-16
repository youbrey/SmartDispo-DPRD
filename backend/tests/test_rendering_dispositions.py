from copy import deepcopy
from pathlib import Path

from docx import Document as WordDocument

from app.services.rendering import _fill_disposition

ROOT = Path(__file__).resolve().parents[2]


def content() -> dict:
    return {
        "sender": "Instansi Pengirim",
        "letter_number": "100/001",
        "letter_date": "2026-09-15",
        "received_date": "2026-09-16",
        "agenda_number": "AG-001",
        "subject": "Koordinasi administrasi",
        "priority": "PENTING",
        "officeholders": {"SEKWAN": {"full_name": "Sekretaris Aktif", "metadata": {}}},
        "dispositions": {
            "SEKWAN": {"directives": ["COORDINATE"], "note": "Koordinasikan segera."}
        },
        "disposition_target_labels_by_actor": {
            "SEKWAN": ["Bagian Umum dan Keuangan", "Komisi I"]
        },
    }


def document_text(word: WordDocument) -> str:
    paragraphs = [paragraph.text for paragraph in word.paragraphs]
    cells = [cell.text for table in word.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs, *cells])


def test_setwan_renderer_prints_dynamic_targets_and_selected_directive() -> None:
    word = WordDocument(ROOT / "templates" / "disposition-setwan-v1.docx")
    _fill_disposition(word, content(), "INCOMING_SECRETARY")
    text = document_text(word)
    assert "☒ Bagian Umum dan Keuangan" in text
    assert "☒ Komisi I" in text
    assert "☒ Koordinasikan" in text


def test_dprd_renderer_prints_sekwan_targets() -> None:
    payload = deepcopy(content())
    payload["dispositions"]["CHAIRMAN"] = {
        "directives": ["FORWARD_COMMISSION_I", "FOLLOW_UP"],
        "note": "Tindak lanjuti.",
    }
    word = WordDocument(ROOT / "templates" / "disposition-dprd-v1.docx")
    _fill_disposition(word, payload, "INCOMING_CHAIRMAN")
    text = document_text(word)
    assert "☒ Bagian Umum dan Keuangan" in text
    assert "☒ Komisi I" in text
    assert "☒ TERUSKAN KE KOMISI I" in text
